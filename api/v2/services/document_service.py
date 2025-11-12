from typing import List
import os
from fastapi import UploadFile
from docx import Document as DocxDocument

from api.v2.models.models import Document, PlaceHolder
from api.v2.repository.document_repository import document_repo_ins
from api.v2.app.langchain.parser import LangChainParser


class DocumentService:
    """Service for handling document upload and parsing with LangChain"""

    def __init__(self):
        self.parser = LangChainParser()

    async def upload_and_parse(self, file: UploadFile) -> Document:
        """Upload a document and parse it to extract placeholders"""

        # Save uploaded file
        upload_dir = "uploads"
        os.makedirs(upload_dir, exist_ok=True)
        file_path = os.path.join(upload_dir, file.filename)

        with open(file_path, "wb") as f:
            content = await file.read()
            f.write(content)

        # Parse document to extract placeholders using LangChain
        placeholders, temp_doc_path = await self.parser.parse_document(file_path)

        # Create document model
        document = Document(
            title=file.filename,
            path=file_path,
            temp_path=temp_doc_path,  # Store temp document path
            placeholders=placeholders,
            analysis_metadata={
                "total_placeholders": len(placeholders),
                "confidence_scores": [
                    p.analysis.confidence_score for p in placeholders if p.analysis
                ],
                "placeholder_types": [
                    p.analysis.inferred_type.value for p in placeholders if p.analysis
                ],
            },
        )

        # Save to database
        print(f"Saving document to database: {document.title}")
        saved_document = await document_repo_ins.save(document)
        print(f"Document saved with ID: {saved_document.id}")

        return saved_document

    async def get_document(self, document_id: str) -> Document:
        """Get a document by ID"""
        document = await document_repo_ins.find_by_id(document_id)
        if not document:
            raise ValueError(f"Document {document_id} not found")
        return document

    async def list_documents(self) -> List[Document]:
        """List all documents"""
        return await document_repo_ins.find_all()


# Singleton instance
document_service = DocumentService()
